"""
scratch/compare_sentences.py
Sentence-by-sentence comparison between marked docx and original manuscripts.
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
import docx
import re
from difflib import SequenceMatcher

root = Path(r"d:\fm_paper_works\hcc-vs-hemangioma")
doc_orig = docx.Document(root / "paper_submission" / "manu_before_revision.docx")
doc_rev1 = docx.Document(root / "paper_submission" / "manuscript_revised.docx")
doc_marked = docx.Document(root / "paper_submission" / "revision2" / "manuscript_revised_marked.docx")

def get_sentences(text):
    # Regex split on sentence endings followed by space
    # Avoid splitting on et al., vs., No., Fig., numbers like 2.90, etc.
    text = re.sub(r'et al\.', 'et al<DOT>', text)
    text = re.sub(r'vs\.', 'vs<DOT>', text)
    text = re.sub(r'No\.', 'No<DOT>', text)
    text = re.sub(r'Fig\.', 'Fig<DOT>', text)
    text = re.sub(r'(\d+)\.(\d+)', r'\1<DECIMAL>\2', text)
    
    splits = re.split(r'([.?!])\s+', text)
    sents = []
    # reconstruct sentence + delimiter
    cur = ""
    for part in splits:
        if part in ['.', '?', '!']:
            cur += part
            sents.append(cur.strip())
            cur = ""
        else:
            cur += part
    if cur.strip():
        sents.append(cur.strip())
        
    restored = []
    for s in sents:
        s = s.replace('et al<DOT>', 'et al.')
        s = s.replace('vs<DOT>', 'vs.')
        s = s.replace('No<DOT>', 'No.')
        s = s.replace('Fig<DOT>', 'Fig.')
        s = s.replace('<DECIMAL>', '.')
        if s.strip():
            restored.append(s.strip())
    return restored

orig_pool = []
for doc in [doc_orig, doc_rev1]:
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            orig_pool.extend(get_sentences(txt))

def clean_for_match(s):
    s = re.sub(r'\[[\d,\s\-]+\]', '', s)
    s = re.sub(r'\([\d,\s\-]+\)', '', s)
    s = re.sub(r'\s+', ' ', s)
    return s.strip()

clean_orig = [clean_for_match(s) for s in orig_pool]
print(f"Original sentences pool size: {len(clean_orig)}")

total_sents = 0
match_count = 0
diff_count = 0

for p_idx, p in enumerate(doc_marked.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    sents = get_sentences(txt)
    print(f"\n=== Paragraph {p_idx} ({len(sents)} sentences) ===")
    for s_idx, s in enumerate(sents):
        total_sents += 1
        cs = clean_for_match(s)
        best_r = 0
        best_orig = ""
        for co, ro in zip(clean_orig, orig_pool):
            if cs == co:
                best_r = 1.0
                best_orig = ro
                break
            r = SequenceMatcher(None, cs, co).ratio()
            if r > best_r:
                best_r = r
                best_orig = ro
        
        # Consider a match if >= 0.82 or exact substring
        is_match = (best_r >= 0.82) or any(cs in co or co in cs for co in clean_orig if len(co) > 20)
        if is_match:
            match_count += 1
            status = "MATCH (BLACK)"
        else:
            diff_count += 1
            status = "REVISED (RED)"
            
        print(f"  [{status} {best_r:.2f}] {s[:60]}...")
        if 0.50 <= best_r < 0.82:
            print(f"      ORIG: {best_orig[:60]}...")

print(f"\nTOTAL: {total_sents} sentences. MATCH(BLACK): {match_count}, REVISED(RED): {diff_count}")
