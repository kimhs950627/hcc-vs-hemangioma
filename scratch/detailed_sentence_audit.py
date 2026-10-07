"""
scratch/detailed_sentence_audit.py
Examine each sentence in generate_revised_manuscript.py and classify as ORIGINAL (BLACK) or REVISED (RED).
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
import docx
import re
from difflib import SequenceMatcher

root = Path(r"d:\fm_paper_works\hcc-vs-hemangioma")
doc_orig = docx.Document(root / "paper_submission" / "manu_before_revision.docx")
doc_rev1 = docx.Document(root / "paper_submission" / "manuscript_revised.docx")

def get_sentences(text):
    text = re.sub(r'et al\.', 'et al<DOT>', text)
    text = re.sub(r'vs\.', 'vs<DOT>', text)
    text = re.sub(r'No\.', 'No<DOT>', text)
    text = re.sub(r'Fig\.', 'Fig<DOT>', text)
    text = re.sub(r'(\d+)\.(\d+)', r'\1<DECIMAL>\2', text)
    
    splits = re.split(r'([.?!])\s+', text)
    sents = []
    cur = ""
    for part in splits:
        if part in ['.', '?', '!']:
            cur += part
            sents.append(cur.strip())
            cur = ""
        else:
            cur += part
    if cur.strip():
        sents.append(cur.strip())
        
    restored = []
    for s in sents:
        s = s.replace('et al<DOT>', 'et al.')
        s = s.replace('vs<DOT>', 'vs.')
        s = s.replace('No<DOT>', 'No.')
        s = s.replace('Fig<DOT>', 'Fig.')
        s = s.replace('<DECIMAL>', '.')
        if s.strip():
            restored.append(s.strip())
    return restored

orig_pool = []
for doc in [doc_orig, doc_rev1]:
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            orig_pool.extend(get_sentences(txt))

def clean_for_match(s):
    s = re.sub(r'\[[\d,\s\-]+\]', '', s)
    s = re.sub(r'\([\d,\s\-]+\)', '', s)
    s = re.sub(r'[\n\r\t]', ' ', s)
    s = re.sub(r'\s+', ' ', s)
    return s.strip()

clean_orig = [clean_for_match(s) for s in orig_pool]

sys.path.append(str(root / "paper_submission" / "revision2"))
import generate_revised_manuscript as grm

print("AUDIT RESULTS:\n" + "="*80)

for sec_idx, (ptype, chunks) in enumerate(grm.doc_structure):
    full_text = "".join(c[0] for c in chunks)
    sents = get_sentences(full_text)
    print(f"\n[{sec_idx}] Type: {ptype} | Sentences: {len(sents)}")
    
    for s_idx, s in enumerate(sents):
        cs = clean_for_match(s)
        best_r = 0
        best_orig = ""
        for co, ro in zip(clean_orig, orig_pool):
            if cs == co:
                best_r = 1.0
                best_orig = ro
                break
            r = SequenceMatcher(None, cs, co).ratio()
            if r > best_r:
                best_r = r
                best_orig = ro
        
        # Exact substring in any original sentence or high ratio
        is_exact = any(cs == co for co in clean_orig)
        is_sub = any(cs in co or co in cs for co in clean_orig if len(cs) > 20 and len(co) > 20)
        
        if is_exact or best_r >= 0.85:
            color = "BLACK (ORIGINAL)"
        elif is_sub or best_r >= 0.70:
            color = "REVIEW (SIMILAR: " + f"{best_r:.2f})"
        else:
            color = "RED (NEW/REVISED)"
            
        print(f"  ({s_idx}) [{color}] {s[:65]}...")
        if "REVIEW" in color:
            print(f"       -> ORIG: {best_orig[:65]}...")
