"""Extract red-marked (revised) runs from the marked manuscript docx."""
import os
import re
import docx

PATH = r"d:\fm_paper_works\hcc-vs-hemangioma\paper_submission\revision2\Manuscript revised marked.docx"
OUT = r"d:\fm_paper_works\hcc-vs-hemangioma\paper_submission\revision2\revised_parts_extracted.md"


def rgb_str(run):
    c = run.font.color
    if c is None:
        return None
    try:
        if c.rgb is None:
            return None
        return str(c.rgb).upper()
    except Exception:
        return None


def is_red(run):
    s = rgb_str(run)
    if not s or len(s) != 6:
        return False
    try:
        r = int(s[0:2], 16); g = int(s[2:4], 16); b = int(s[4:6], 16)
    except Exception:
        return False
    return r >= 140 and g <= 110 and b <= 110


def main():
    doc = docx.Document(PATH)
    lines = []
    n_red_par = 0
    n_par = 0
    for i, p in enumerate(doc.paragraphs):
        n_par += 1
        text = p.text.strip()
        style = p.style.name if p.style is not None else ""
        if not text:
            continue
        reds = [r.text for r in p.runs if is_red(r) and r.text.strip()]
        if reds:
            n_red_par += 1
            red_clean = re.sub(r"\s+", " ", "".join(reds)).strip()
            full_clean = re.sub(r"\s+", " ", "".join(
                (r.text if not is_red(r) else f"«{r.text}»") for r in p.runs
            )).strip()
            lines.append(f"### [p{i}]")
            lines.append(f"RED: {red_clean}")
            lines.append(f"CTX: {full_clean}")
            lines.append("")
        else:
            if len(text) < 90 and (style.lower().startswith("heading") or text[:1].isdigit()
                                   or text[:1] in ("1", "2", "3", "4", "5")
                                   or text.startswith(("Abstract", "Table", "초록"))):
                lines.append(f"\n## [p{i}] {text}\n")
    total_red = sum(1 for p in doc.paragraphs for r in p.runs if is_red(r) and r.text.strip())
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(f"# Revised (red) parts — {os.path.basename(PATH)}\n\n")
        f.write(f"paragraphs={n_par}, paragraphs_with_red={n_red_par}, red_runs={total_red}\n")
        f.write("(« » = red/revised segment inside the paragraph)\n\n")
        f.write("\n".join(lines))
    print(f"paragraphs={n_par} paragraphs_with_red={n_red_par} red_runs={total_red}")
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
