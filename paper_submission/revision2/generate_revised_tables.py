"""
Generate revised Tables and Figure Captions in Word format (docx) for KJFP revision.
File: tables_and_figures_revised.docx
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"d:\fm_paper_works\hcc-vs-hemangioma\paper_submission\revision2"
os.makedirs(OUT_DIR, exist_ok=True)

doc = docx.Document()
for s in doc.sections:
    s.top_margin = Inches(1.0)
    s.bottom_margin = Inches(1.0)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)

RED = RGBColor(220, 20, 20)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_table_header(table, headers):
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(9.5)
        set_cell_background(hdr_cells[i], "F0F0F0")

def add_caption(doc, text, is_revised=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.bold = True
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(11)
    if is_revised:
        run.font.color.rgb = RED

def add_footnote(doc, text, is_revised=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(8.5)
    run.font.italic = True
    if is_revised:
        run.font.color.rgb = RED

# --- Table 1 ---
add_caption(doc, "Table 1. Patient and Lesion Demographics of Full Cohort (SMC-LUD, n=1,021)", is_revised=True)
t1_data = [
    ["항목 (Characteristics)", "HCC (환자 n=600)", "Hemangioma (환자 n=421)"],
    ["나이, 평균 ± 표준편차 (범위), 세", "66.65 ± 11.02 (20–90)", "55.50 ± 12.76 (30–90)"],
    ["남성, n (%)", "491 (81.8%)", "174 (41.3%)"],
    ["여성, n (%)", "109 (18.2%)", "247 (58.7%)"],
    ["키, 평균 ± 표준편차, cm", "165.74 ± 16.39", "163.72 ± 8.45"],
    ["몸무게, 평균 ± 표준편차 (범위), kg", "67.44 ± 11.24 (42–104)", "63.85 ± 12.95 (34–118)"],
    ["병변 크기 (최대 직경), 중앙값 / Q1 / Q3, cm", "2.90 / 2.10 / 4.50", "—"],
    ["UICC 병기 (2000), n (%)", "", ""],
    ["  pT1", "58 (9.7%)", "—"],
    ["  pT2", "120 (20.0%)", "—"],
    ["  pT3", "353 (58.8%)", "—"],
    ["  pT4", "69 (11.5%)", "—"],
    ["AJCC 병기 (2010), n (%)", "", ""],
    ["  pT1", "164 (27.3%)", "—"],
    ["  pT2", "393 (65.5%)", "—"],
    ["  pT3", "22 (3.7%)", "—"],
    ["  pT4", "21 (3.5%)", "—"],
]

t1 = doc.add_table(rows=len(t1_data), cols=3)
t1.style = 'Table Grid'
add_table_header(t1, t1_data[0])
for r_idx, row in enumerate(t1_data[1:], start=1):
    for c_idx, val in enumerate(row):
        cell = t1.rows[r_idx].cells[c_idx]
        cell.text = val
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(9)
add_footnote(doc, "Data source: Tak et al. Sci Data (2026)[13]. Values represent full SMC-LUD cohort (n=1,021 patients). Note that tumor diameter and pathological stages were only documented for HCC cases in the original public repository.", is_revised=True)

# --- Table 1B ---
add_caption(doc, "Table 1B. Dataset Allocation for Clean Subset (744 Patients, 2,656 Images)", is_revised=False)
t1b_data = [
    ["Diagnosis", "No. of Patients", "No. of Images", "Train Set", "Validation Set", "Test Set"],
    ["Hemangioma", "323", "1,267", "886", "253", "128"],
    ["HCC", "421", "1,389", "972", "277", "140"],
    ["Total (Clean)", "744", "2,656", "1,858", "530", "268"]
]
t1b = doc.add_table(rows=len(t1b_data), cols=6)
t1b.style = 'Table Grid'
add_table_header(t1b, t1b_data[0])
for r_idx, row in enumerate(t1b_data[1:], start=1):
    for c_idx, val in enumerate(row):
        cell = t1b.rows[r_idx].cells[c_idx]
        cell.text = val
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(9)
            if r_idx == len(t1b_data)-1:
                run.font.bold = True
add_footnote(doc, "Values follow the patient-level split published by Tak et al. (2026). Images with calipers or measuring overlays were excluded to prevent shortcut feature learning.", is_revised=True)

# --- Table 2 ---
add_caption(doc, "Table 2. Ablation: Backbone × Training Mode (Validation Set, Confidence Score)", is_revised=True)
t2_data = [
    ["#", "Backbone", "Training Mode", "Cutoff", "AUROC", "Sensitivity (%)", "Specificity (%)", "F1 Score"],
    ["1", "ResNet50V2", "CE Only", "0.0002", "1.000", "100.00", "100.00", "1.0000"],
    ["2", "EfficientNetV2B0", "CE Only", "0.0016", "0.9999", "100.00", "99.60", "0.9982"],
    ["3", "ResNet50V2", "CE + SupCon", "0.0002", "1.000", "99.64", "100.00", "0.9982"],
    ["4", "EfficientNetV2B0*", "CE + SupCon", "0.0011", "1.000", "100.00", "100.00", "1.0000"],
    ["5", "ResNet50V2", "NNCLR (SSL)", "0.0003", "0.9994", "99.64", "99.21", "0.9946"],
    ["6", "EfficientNetV2B0", "NNCLR (SSL)", "0.0038", "0.9986", "98.19", "98.42", "0.9837"],
]
t2 = doc.add_table(rows=len(t2_data), cols=8)
t2.style = 'Table Grid'
add_table_header(t2, t2_data[0])
for r_idx, row in enumerate(t2_data[1:], start=1):
    for c_idx, val in enumerate(row):
        cell = t2.rows[r_idx].cells[c_idx]
        cell.text = val
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8.5)
            if r_idx == 4: # primary model
                run.font.bold = True
add_footnote(doc, "*Primary model (#4). Cutoffs were determined on the validation set using Youden's J statistic. All values are recomputed from raw prediction logs. CE, Cross-Entropy; SupCon, Supervised Contrastive Learning; NNCLR, Nearest-Neighbor Contrastive Learning (2-stage: SSL pre-training followed by CE+SupCon fine-tuning).", is_revised=True)

# --- Table 3 ---
add_caption(doc, "Table 3. Primary Model Classification Performance (EfficientNetV2B0 + CE+SupCon)", is_revised=True)
t3_data = [
    ["Metric", "Validation Set (n=530)", "Test Set (n=268)"],
    ["AUROC", "1.000", "1.000"],
    ["Accuracy (%) (95% CI)", "100.00 (99.31–100.00)", "99.63 (97.94–99.99)"],
    ["Sensitivity (%) (95% CI)", "100.00 (98.68–100.00)", "100.00 (97.40–100.00)"],
    ["Specificity (%) (95% CI)", "100.00 (98.55–100.00)", "99.22 (95.72–99.98)"],
    ["Positive Predictive Value (%)", "100.00", "99.29"],
    ["Negative Predictive Value (%)", "100.00", "100.00"],
    ["F1 Score", "1.0000", "0.9964"],
    ["Confusion Matrix (TP / FP / FN / TN)", "277 / 0 / 0 / 253", "140 / 1 / 0 / 127"]
]
t3 = doc.add_table(rows=len(t3_data), cols=3)
t3.style = 'Table Grid'
add_table_header(t3, t3_data[0])
for r_idx, row in enumerate(t3_data[1:], start=1):
    for c_idx, val in enumerate(row):
        cell = t3.rows[r_idx].cells[c_idx]
        cell.text = val
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(9)
add_footnote(doc, "Operating threshold (cutoff = 0.00109) was fixed on the validation set using Youden's J statistic and applied unchanged to the independent test set (n=268; HCC 140, hemangioma 128). Test set: AUROC 1.000; accuracy 99.63% (95% CI 97.94–99.99); sensitivity 100.00% (95% CI 97.40–100.00); specificity 99.22% (95% CI 95.72–99.98); PPV 99.29%; NPV 100.00%; F1 0.9964. 95% confidence intervals (CI) are exact Clopper-Pearson binomial intervals estimated at the image level. In the test set, misclassification occurred in exactly 1 hemangioma image (FP=1, FN=0).", is_revised=True)

# --- Table 4 ---
add_caption(doc, "Table 4. Comparison of Dual-Output Metrics in Primary Model (EfficientNetV2B0 + CE+SupCon)", is_revised=True)
t4_data = [
    ["Output", "Score Type", "Prototype Type", "Cutoff (Val)", "AUROC (Val)", "AUROC (Test)", "Test Sens (%)", "Test Spec (%)", "Test Errors (FP / FN)"],
    ["ROC-A", "Confidence Score", "—", "0.0011", "1.000", "1.000", "100.00", "99.22", "1 / 0"],
    ["ROC-B", "HCC Cosine Score", "Mean (n=1)", "-0.0048", "1.000", "1.000", "100.00", "99.22", "1 / 0"],
    ["ROC-C", "Δscore", "Mean (n=1)", "-0.9585", "1.000", "1.000", "100.00", "99.22", "1 / 0"],
    ["ROC-B", "HCC Cosine Score", "k-means (k=4)", "0.0288", "1.000", "1.000", "100.00", "99.22", "1 / 0"],
    ["ROC-C", "Δscore", "k-means (k=4)", "-0.9265", "1.000", "1.000", "100.00", "99.22", "1 / 0"]
]
t4 = doc.add_table(rows=len(t4_data), cols=9)
t4.style = 'Table Grid'
add_table_header(t4, t4_data[0])
for r_idx, row in enumerate(t4_data[1:], start=1):
    for c_idx, val in enumerate(row):
        cell = t4.rows[r_idx].cells[c_idx]
        cell.text = val
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8)
add_footnote(doc, "All prototypes were derived strictly from the training set. Cutoffs were determined on the validation set. Note: In the test set (n=268), exactly the same single hemangioma image was misclassified as false positive across all 5 output definitions; thus, no statistical non-inferiority or incremental triage benefit could be confirmed.", is_revised=True)

# --- Table 4B ---
add_caption(doc, "Table 4B. Test Set AUROC Comparison across All Evaluated Models", is_revised=True)
t4b_data = [
    ["Backbone", "Training Mode", "Confidence (A)", "Cosine-B (Mean)", "Cosine-B (k-means)", "Δscore-C (Mean)", "Δscore-C (k-means)"],
    ["EfficientNetV2B0", "CE Only", "0.9999", "0.9999", "1.0000", "0.9999", "0.9999"],
    ["EfficientNetV2B0*", "CE + SupCon", "1.0000", "1.0000", "1.0000", "1.0000", "1.0000"],
    ["EfficientNetV2B0", "NNCLR (SSL)", "1.0000", "0.8930", "0.9763", "0.9996", "0.9999"],
    ["ResNet50V2", "CE Only", "1.0000", "1.0000", "1.0000", "1.0000", "1.0000"],
    ["ResNet50V2", "CE + SupCon", "0.9999", "0.9999", "0.9999", "0.9999", "0.9999"],
    ["ResNet50V2", "NNCLR (SSL)", "0.9994", "0.0684", "0.1235", "0.9891", "0.9984"]
]
t4b = doc.add_table(rows=len(t4b_data), cols=7)
t4b.style = 'Table Grid'
add_table_header(t4b, t4b_data[0])
for r_idx, row in enumerate(t4b_data[1:], start=1):
    for c_idx, val in enumerate(row):
        cell = t4b.rows[r_idx].cells[c_idx]
        cell.text = val
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Malgun Gothic"
            run.font.size = Pt(8.5)
            if r_idx == 2:
                run.font.bold = True
add_footnote(doc, "*Primary model. AUROC values are re-evaluated on the independent test set (n=268). In ResNet50V2 + NNCLR, the AUROC of single prototype HCC cosine score collapsed to 0.0684, representing label-space inversion resulting from unguided self-supervised pre-training representations.", is_revised=True)

# --- Figure Captions ---
p_fig = doc.add_paragraph()
p_fig.paragraph_format.space_before = Pt(20)
run_fig = p_fig.add_run("그림 설명 (Figure Captions)")
run_fig.font.bold = True
run_fig.font.name = "Malgun Gothic"
run_fig.font.size = Pt(13)

captions = [
    ("Figure 1. t-SNE Embedding Space Visualization (Primary Model: EfficientNetV2B0 + CE+SupCon).",
     "Embedding distribution of training and validation sets showing clear intra-class compactness and inter-class separation between HCC (red) and hemangioma (teal) clusters. Left panel illustrates single mean prototype centres; right panel depicts k-means cluster prototypes (k=4).",
     True),
    ("Figure 2. Confusion Matrices of Primary Model on Held-out Test Set (n=268).",
     "Across all three primary outputs—(A) Confidence Score, (B) HCC Cosine Score (Mean), and (C) Δscore (Mean)—classification on the test set yielded identical performance: 140 true positives, 127 true negatives, 1 false positive, and 0 false negatives (sensitivity 100.0%, specificity 99.22%).",
     True),
    ("Figure 3. Triple Receiver Operating Characteristic (ROC) Curves on Test Set.",
     "ROC curves for ROC-A (Confidence Score), ROC-B (HCC Cosine Score), and ROC-C (Δscore). All three metrics achieve an AUROC of 1.000 on the held-out test set under the clean subset benchmark.",
     False),
    ("Figure 4. HCC Cosine Score Distributions with and without Supervised Contrastive Learning (SupCon) in True HCC Cases (n=417; Validation 277 + Test 140).",
     "The same held-out HCC images were scored by the CE-only and CE+SupCon models (HCC cosine score to the training-set mean HCC prototype). (A) Density distribution, EfficientNetV2B0. (B) Paired individual shift, EfficientNetV2B0: mean 0.783 ± 0.256 (median 0.882) to 0.908 ± 0.160 (median 0.970); mean paired difference +0.125 (95% CI 0.110–0.140); paired t = 16.94; Wilcoxon signed-rank p < 0.001; score increased in 94.7% of images; variance 0.065 to 0.026 (−60.7%; Pitman-Morgan t = 16.84, p < 0.001). (C) Density distribution, ResNet50V2. (D) Paired individual shift, ResNet50V2: mean 0.774 ± 0.265 to 0.878 ± 0.230; mean paired difference +0.104; paired t = 10.98; Wilcoxon p < 0.001; variance 0.070 to 0.053 (−24.3%; Pitman-Morgan t = 4.06, p < 0.001). CE, cross-entropy.",
     True),
    ("Figure 5. Hemangioma Cosine Score Distributions with and without Supervised Contrastive Learning (SupCon) in True Hemangioma Cases (n=381; Validation 253 + Test 128).",
     "The same held-out hemangioma images were scored by the CE-only and CE+SupCon models (hemangioma cosine score to the training-set mean hemangioma prototype). (A) Density distribution, EfficientNetV2B0. (B) Paired individual shift, EfficientNetV2B0: mean 0.937 ± 0.050 (median 0.955) to 0.985 ± 0.006 (median 0.985); mean paired difference +0.048 (95% CI 0.043–0.052); paired t = 19.71; Wilcoxon signed-rank p < 0.001; variance 0.00253 to 0.00003 (Pitman-Morgan t = 106.23, p < 0.001). (C) Density distribution, ResNet50V2. (D) Paired individual shift, ResNet50V2: mean 0.966 ± 0.019 to 0.990 ± 0.004; mean paired difference +0.024 (95% CI 0.022–0.026); paired t = 26.54; Wilcoxon p < 0.001; variance 0.00036 to 0.00002 (Pitman-Morgan t = 49.18, p < 0.001). The large relative variance reduction partly reflects the already small baseline variance under CE-only training.",
     True)
]

for title, desc, is_rev in captions:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(title + "\n")
    r1.font.bold = True
    r1.font.name = "Malgun Gothic"
    r1.font.size = Pt(10)
    if is_rev:
        r1.font.color.rgb = RED
        
    r2 = p.add_run(desc)
    r2.font.name = "Malgun Gothic"
    r2.font.size = Pt(9.5)
    if is_rev:
        r2.font.color.rgb = RED

doc_path = os.path.join(OUT_DIR, "tables_and_figures_revised.docx")
doc.save(doc_path)
print(f"Saved revised tables and figure captions to: {doc_path}")
