#!/usr/bin/env python3
"""
Manuscript revision script for manu_before_revision.docx
─────────────────────────────────────────────────────────
1. Korean abstract ≤ 1,000 chars (with spaces), remove refs & sub-headers
2. Renumber body references to sequential first-appearance order
3. Reorder reference list, remove uncited ref [11]
4. Insert English abstract + keywords (no sub-headers) after Korean keywords
5. Save as manuscript_revised.docx + manuscript_revised.txt
"""

import re
import sys
from copy import deepcopy
from docx import Document
from lxml import etree

# ════════════════════════════════════════════════════════════
# Configuration
# ════════════════════════════════════════════════════════════

INPUT  = 'manu_before_revision.docx'
OUT_DOCX = 'manuscript_revised.docx'
OUT_TXT  = 'manuscript_revised.txt'

W   = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
XSP = '{http://www.w3.org/XML/1998/namespace}space'

# ── Reference first-appearance order in body ──
FIRST_APP = [1,2,3,4,21,22,23,24,5,6,16,12,9,8,7,13,10,14,20,17,18,19,15]
OLD2NEW = {old: new for new, old in enumerate(FIRST_APP, 1)}
NEW2OLD = {v: k for k, v in OLD2NEW.items()}

# ── Shortened Korean abstract content (4 paragraphs, NO sub-headers) ──
KOR_BG = (
    "B-mode 초음파는 간세포암(HCC) 감시의 핵심 도구이나 민감도가 제한적이며, "
    "기존 딥러닝 모델은 단일 softmax 출력만을 제공하여 "
    "임상의의 유사성 기반 추론에 대응하지 못한다."
)
KOR_METHODS = (
    "공개 데이터셋 SMC-LUD의 Clean subset"
    "(744명, 2,656장)을 이용한 후향적 연구다. "
    "EfficientNetV2B0 기반 hybrid vision transformer를 "
    "교차 엔트로피와 supervised contrastive learning(SupCon) "
    "결합 학습하였다. Softmax를 confidence score로, "
    "임베딩과 HCC prototype 간 cosine similarity를 HCC cosine score, "
    "두 클래스 간 마진을 \u0394score로 정의하는 이중 출력 체계를 "
    "구성하였다(IRB No. PMC 2026-07-001)."
)
KOR_RESULTS = (
    "주력 모델의 confidence score AUROC는 검증 및 테스트 세트 "
    "모두 1.000이었으며, 테스트 세트에서 민감도 100.0%, "
    "특이도 99.2%를 달성하였다. HCC cosine score와 \u0394score "
    "AUROC 1.000으로 confidence score와 동등하였다"
    "(DeLong p=1.000). SupCon 적용 시 true HCC cases에서 "
    "HCC cosine score가 유의하게 상승하였다(p<0.001). "
    "NNCLR 사전학습 후 미세조정한 2단계 학습은 "
    "CE+SupCon 대비 분류 성능이 열등하였다."
)
KOR_CONCL = (
    "본 연구는 B-mode 초음파에서 임상의의 유사성 기반 추론에 대응하는 "
    "연속형 영상표지자 후보(HCC cosine score, \u0394score)를 제안하였으며, "
    "SupCon 학습이 cosine 기반 표지자의 임상적 타당성 확보에 "
    "필수적임을 확인하였다."
)
KOR_KW = "핵심어: 간세포암; 초음파; 영상표지자; cosine similarity; supervised contrastive learning; 이중 출력; 임상 의사결정 지원"

# ── English abstract (direct translation, NO sub-headers) ──
ENG_BG = (
    "B-mode ultrasonography is a core tool for hepatocellular carcinoma "
    "(HCC) surveillance; however, its sensitivity remains limited. "
    "Existing deep learning models provide only a single softmax output, "
    "which does not correspond to the similarity-based reasoning "
    "performed by clinicians."
)
ENG_METHODS = (
    "This was a retrospective study using the Clean subset "
    "(744 patients, 2,656 images) from the publicly available "
    "Samsung Medical Center\u2013Liver Ultrasound Dataset (SMC-LUD). "
    "A hybrid vision transformer based on EfficientNetV2B0 was trained "
    "with a combination of cross-entropy and supervised contrastive "
    "learning (SupCon). A dual-output system was constructed in which "
    "the softmax output was defined as the confidence score, "
    "the cosine similarity between the embedding and the HCC prototype "
    "as the HCC cosine score, and the margin between the two class "
    "prototypes as the \u0394score (IRB No. PMC 2026-07-001)."
)
ENG_RESULTS = (
    "The AUROC of the confidence score of the primary model was 1.000 "
    "in both the validation and test sets, achieving a sensitivity of "
    "100.0% and a specificity of 99.2% on the test set. "
    "The HCC cosine score and \u0394score also demonstrated an AUROC of "
    "1.000, equivalent to the confidence score (DeLong p = 1.000). "
    "With SupCon applied, the HCC cosine score in true HCC cases was "
    "significantly elevated (p < 0.001). The two-stage training with "
    "NNCLR pre-training followed by fine-tuning was inferior in "
    "classification performance compared to CE+SupCon alone."
)
ENG_CONCL = (
    "This study proposed novel continuous imaging biomarker candidates "
    "(HCC cosine score, \u0394score) that correspond to similarity-based "
    "clinical reasoning on B-mode ultrasonography, and confirmed that "
    "SupCon training is essential for ensuring the clinical validity "
    "of cosine-based biomarkers."
)
ENG_KW = (
    "Keywords: hepatocellular carcinoma; ultrasonography; "
    "imaging biomarker; cosine similarity; supervised contrastive "
    "learning; dual output; clinical decision support"
)

# ════════════════════════════════════════════════════════════
# Helper functions
# ════════════════════════════════════════════════════════════

def set_text(para, text):
    """Replace all text in paragraph, keeping first run's char format."""
    if para.runs:
        para.runs[0].text = text
        for r in list(para.runs[1:]):
            r._element.getparent().remove(r._element)
    else:
        para.add_run(text)


def ref_sub(m):
    """Regex callback: [old] -> [new]."""
    n = int(m.group(1))
    return f'[{OLD2NEW[n]}]' if n in OLD2NEW else m.group(0)


def renumber_para(para):
    """Renumber [N] citations in a single paragraph."""
    full = para.text
    if not re.search(r'\[\d+\]', full):
        return False
    changed = False
    for run in para.runs:
        if run.text and re.search(r'\[\d+\]', run.text):
            new_t = re.sub(r'\[(\d+)\]', ref_sub, run.text)
            if new_t != run.text:
                run.text = new_t
                changed = True
    if not changed:
        new_full = re.sub(r'\[(\d+)\]', ref_sub, full)
        if new_full != full:
            set_text(para, new_full)
            changed = True
    return changed


def make_p(text, fmt_para):
    """Create a new <w:p> element with formatting cloned from fmt_para."""
    p = deepcopy(fmt_para._element)
    for ch in list(p):
        if ch.tag != f'{W}pPr':
            p.remove(ch)
    if text:
        r = etree.SubElement(p, f'{W}r')
        if fmt_para.runs:
            rpr = fmt_para.runs[0]._element.find(f'{W}rPr')
            if rpr is not None:
                r.insert(0, deepcopy(rpr))
        t = etree.SubElement(r, f'{W}t')
        t.text = text
        t.set(XSP, 'preserve')
    return p


# ════════════════════════════════════════════════════════════
# Main
# ════════════════════════════════════════════════════════════

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    doc = Document(INPUT)
    P = doc.paragraphs
    body_xml = doc.element.body
    print(f'Loaded {INPUT} -- {len(P)} paragraphs\n')

    # ────────────────────────────────────────────────────────
    # Step 1  Korean abstract: remove sub-headers, shorten content,
    #         remove refs. Original paras [3]-[17]:
    #   [3] Abstract (header - KEEP)
    #   [4] empty (KEEP)
    #   [5] 연구 배경 (Background)  ← DELETE
    #   [6] content                 ← REPLACE
    #   [7] empty                   ← DELETE
    #   [8] 방법 (Methods)          ← DELETE
    #   [9] content                 ← REPLACE
    #  [10] empty                   ← DELETE
    #  [11] 결과 (Results)          ← DELETE
    #  [12] content                 ← REPLACE
    #  [13] empty                   ← DELETE
    #  [14] 결론 (Conclusions)      ← DELETE
    #  [15] content                 ← REPLACE
    #  [16] empty                   ← DELETE
    #  [17] 핵심어: ...             ← KEEP (unchanged)
    # ────────────────────────────────────────────────────────

    # First: update content paragraphs with shortened text
    set_text(P[6],  KOR_BG)
    set_text(P[9],  KOR_METHODS)
    set_text(P[12], KOR_RESULTS)
    set_text(P[15], KOR_CONCL)

    # Delete sub-header and empty paragraphs (in reverse to preserve indices)
    # Paragraphs to delete: 16(empty), 14(결론), 13(empty), 11(결과), 
    #                        10(empty), 8(방법), 7(empty), 5(연구배경)
    for idx in [16, 14, 13, 11, 10, 8, 7, 5]:
        body_xml.remove(P[idx]._element)

    # After deletion, re-read paragraphs
    P = doc.paragraphs
    print(f'After abstract cleanup: {len(P)} paragraphs')

    # Now abstract is: [3] Abstract, [4] empty, [5] bg, [6] methods,
    #                  [7] results, [8] conclusions, [9] keywords
    # Calculate char count for abstract (paras 5-9)
    abs_text = ''.join(P[i].text for i in range(5, 10))
    abs_chars = len(abs_text)
    abs_refs = re.findall(r'\[\d+\]', abs_text)
    status1 = 'OK' if abs_chars <= 1000 else f'OVER by {abs_chars - 1000}'
    print(f'[Step 1] Korean abstract: {abs_chars} chars -- {status1}')
    print(f'         Refs in abstract: {abs_refs if abs_refs else "none (OK)"}')

    # ────────────────────────────────────────────────────────
    # Step 2  Renumber body references
    # ────────────────────────────────────────────────────────
    # Find body range: from "1. 서론" to before "참고문헌"
    intro_idx = None
    ref_hdr_idx = None
    for i in range(len(P)):
        if '1. 서론' in P[i].text and intro_idx is None:
            intro_idx = i
        if '참고문헌' in P[i].text and intro_idx is not None:
            ref_hdr_idx = i
            break

    print(f'\n  Body range: paras [{intro_idx}] to [{ref_hdr_idx - 1}]')

    cross_runs = []
    for i in range(intro_idx, ref_hdr_idx):
        full_before = P[i].text
        if not re.search(r'\[\d+\]', full_before):
            continue
        any_run = any(re.search(r'\[\d+\]', r.text) for r in P[i].runs if r.text)
        if not any_run and re.search(r'\[\d+\]', full_before):
            cross_runs.append(i)
        renumber_para(P[i])

    # Verify sequential
    body_txt = ' '.join(P[i].text for i in range(intro_idx, ref_hdr_idx))
    seen, order = set(), []
    for r in (int(x) for x in re.findall(r'\[(\d+)\]', body_txt)):
        if r not in seen:
            seen.add(r); order.append(r)
    seq = all(order[j] <= order[j+1] for j in range(len(order)-1))
    print(f'[Step 2] Body ref order: {order}')
    print(f'         Sequential: {seq}')
    if cross_runs:
        print(f'         Cross-run consolidations: paras {cross_runs}')

    # ────────────────────────────────────────────────────────
    # Step 3  Reorder reference list
    # ────────────────────────────────────────────────────────
    # Find first reference paragraph (starts with "1. ")
    ref_start = None
    for i in range(ref_hdr_idx, len(P)):
        if P[i].text.strip().startswith('1.'):
            ref_start = i
            break

    print(f'\n  Reference list starts at para [{ref_start}]')

    # Deep-copy all 24 reference elements
    ref_elems = {}
    for n in range(1, 25):
        ref_elems[n] = deepcopy(P[ref_start + n - 1]._element)

    # Remove originals
    for n in range(1, 25):
        body_xml.remove(P[ref_start + n - 1]._element)

    # Find anchor (empty line before ref list = ref_start - 1)
    P = doc.paragraphs  # re-read after removal
    # The anchor is the paragraph right before where refs were
    # After removal, we need to find the "참고문헌" header or the empty line after it
    anchor = None
    for i in range(len(P)):
        if '참고문헌' in P[i].text:
            # Anchor is the next paragraph (empty line) or this one
            if i + 1 < len(P):
                anchor = P[i + 1]._element if P[i + 1].text.strip() == '' else P[i]._element
            else:
                anchor = P[i]._element
            break

    # Re-insert in new order (reverse iteration)
    for new_num in range(23, 0, -1):
        old_num = NEW2OLD[new_num]
        elem = ref_elems[old_num]
        # Fix leading number
        for t_el in elem.iter(f'{W}t'):
            if t_el.text and re.match(r'^\d+\.', t_el.text):
                t_el.text = re.sub(r'^\d+\.', f'{new_num}.', t_el.text)
                break
        anchor.addnext(elem)

    print(f'[Step 3] Reference list reordered (23 refs; old #11 removed)')

    # ────────────────────────────────────────────────────────
    # Step 4  Insert English abstract after Korean keywords
    #         NO sub-headers, just content paragraphs
    # ────────────────────────────────────────────────────────
    P = doc.paragraphs  # re-read

    # Find Korean keywords paragraph (핵심어)
    kw_idx = None
    for i in range(len(P)):
        if P[i].text.startswith('핵심어:'):
            kw_idx = i
            break

    kw_elem = P[kw_idx]._element
    cnt_fmt = P[kw_idx - 1]  # content paragraph format (conclusions para)
    emp_fmt_idx = kw_idx + 1  # empty paragraph after keywords
    if emp_fmt_idx < len(P) and P[emp_fmt_idx].text.strip() == '':
        emp_fmt = P[emp_fmt_idx]
    else:
        emp_fmt = P[4]  # fallback: empty para before abstract
    kw_fmt = P[kw_idx]
    abs_hdr_fmt = P[3]  # "Abstract" header style

    # Items listed from LAST -> FIRST (addnext reverses order)
    eng_items = [
        (ENG_KW,              kw_fmt),        # English Keywords
        ('',                  emp_fmt),
        (ENG_CONCL,           cnt_fmt),       # Conclusions content
        ('',                  emp_fmt),
        (ENG_RESULTS,         cnt_fmt),       # Results content
        ('',                  emp_fmt),
        (ENG_METHODS,         cnt_fmt),       # Methods content
        ('',                  emp_fmt),
        (ENG_BG,              cnt_fmt),       # Background content
        ('',                  emp_fmt),
        ('English Abstract',  abs_hdr_fmt),   # Header only
        ('',                  emp_fmt),        # Separator
    ]

    for text, fmt in eng_items:
        kw_elem.addnext(make_p(text, fmt))

    print(f'[Step 4] English abstract inserted ({len(eng_items)} paragraphs, no sub-headers)')

    # ────────────────────────────────────────────────────────
    # Step 5  Save
    # ────────────────────────────────────────────────────────
    doc.save(OUT_DOCX)
    print(f'\n==> Saved: {OUT_DOCX}')

    # Re-read and export TXT
    doc2 = Document(OUT_DOCX)
    with open(OUT_TXT, 'w', encoding='utf-8') as f:
        for i, p in enumerate(doc2.paragraphs):
            f.write(f'[{i}] {p.text}\n')
    print(f'==> Saved: {OUT_TXT}')

    # ────────────────────────────────────────────────────────
    # Step 6  Final verification
    # ────────────────────────────────────────────────────────
    print('\n===== Final Verification =====')
    P2 = doc2.paragraphs
    total = len(P2)
    print(f'Total paragraphs: {total}')

    # Find abstract range in new doc
    abs_start = abs_end = None
    for i in range(total):
        if P2[i].text == 'Abstract':
            abs_start = i + 2  # skip Abstract header and empty line
        if P2[i].text.startswith('핵심어:'):
            abs_end = i + 1  # include keywords
            break

    if abs_start and abs_end:
        abs2_text = ''.join(P2[i].text for i in range(abs_start, abs_end))
        abs2_chars = len(abs2_text)
        abs2_refs = re.findall(r'\[\d+\]', abs2_text)
        print(f'Korean abstract chars: {abs2_chars}  {"<=1000 OK" if abs2_chars <= 1000 else "EXCEEDED"}')
        print(f'Abstract refs: {abs2_refs if abs2_refs else "none (OK)"}')

    # Check English abstract
    eng_found = any('English Abstract' in P2[i].text for i in range(total))
    print(f'English abstract present: {eng_found}')

    # Verify body ref order
    intro2 = ref_hdr2 = None
    for i in range(total):
        if '서론' in P2[i].text and intro2 is None:
            intro2 = i
        if '참고문헌' in P2[i].text and intro2 is not None:
            ref_hdr2 = i
            break
    if intro2 and ref_hdr2:
        body2 = ' '.join(P2[i].text for i in range(intro2, ref_hdr2))
        seen2, order2 = set(), []
        for r in (int(x) for x in re.findall(r'\[(\d+)\]', body2)):
            if r not in seen2:
                seen2.add(r); order2.append(r)
        seq2 = all(order2[j] <= order2[j+1] for j in range(len(order2)-1))
        print(f'Body ref first-appearance: {order2}')
        print(f'Sequential: {seq2}')

    # Print ref list for verification
    print('\n--- Reference list (first 5) ---')
    ref_start2 = None
    for i in range(total):
        if P2[i].text.strip().startswith('1.') and '참고문헌' not in P2[i].text:
            if any('참고문헌' in P2[j].text for j in range(max(0, i-3), i)):
                ref_start2 = i
                break
    if ref_start2:
        for i in range(ref_start2, min(ref_start2 + 5, total)):
            print(f'  [{i}] {P2[i].text[:80]}...')

    print('\nDone.')


if __name__ == '__main__':
    main()
