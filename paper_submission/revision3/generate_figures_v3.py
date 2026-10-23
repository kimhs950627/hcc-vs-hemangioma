"""Build figures_v3.docx: every figure with its caption, one figure per page.

Captions come from captions_v3.py, the same source used by tables_v3.docx, so
the two documents cannot drift apart. Images are the delivery files in
figures/ (Figure 2 has three panels; the validation ROC panel is not included).
"""
import os
import pathlib

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
OUT_DIR = str(HERE)
FIG = HERE / "figures"
RED = RGBColor(220, 20, 20)
GRAY = RGBColor(110, 110, 110)

from captions_v3 import CAPTIONS

# figure -> [(panel label or None, file, width inches)]
LAYOUT = {
    1: [(None, "fig 1 TSNE visualization.jpg", 6.3)],
    2: [("A", "fig 2-A Confusion matrix-A.jpg", 4.4),
        ("B", "fig 2-B Confusion matrix-B.jpg", 3.9),
        ("C", "fig 2-C Confusion matrix-C.jpg", 3.9)],
    3: [(None, "fig 3.jpg", 5.5)],
    4: [(None, "fig_4_hcc_cosine_distribution_comparison.png", 6.3)],
    5: [(None, "fig_5_hemangioma_cosine_distribution_comparison.png", 6.3)],
}

doc = docx.Document()
for s in doc.sections:
    s.top_margin = docx.shared.Inches(0.9)
    s.bottom_margin = docx.shared.Inches(0.9)
    s.left_margin = docx.shared.Inches(1.0)
    s.right_margin = docx.shared.Inches(1.0)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Figures")
r.font.name = "Malgun Gothic"
r.font.size = Pt(16)
r.font.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("HCC vs hemangioma on B-mode ultrasound — revision 3")
r.font.name = "Malgun Gothic"
r.font.size = Pt(11)
r.font.italic = True

note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = note.add_run("Performance figures are test set only (n=268). Cutoffs and their equivalence "
                 "intervals are derived on the validation set.")
r.font.name = "Malgun Gothic"
r.font.size = Pt(9)
r.font.color.rgb = GRAY

missing = []
for idx, (cap_title, cap_desc, is_rev) in enumerate(CAPTIONS, start=1):
    panels = LAYOUT.get(idx, [])
    if idx > 1:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for label, fname, width in panels:
        fpath = FIG / fname
        if not fpath.exists():
            missing.append(fname)
            continue
        if label:
            lp = doc.add_paragraph()
            lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lp.paragraph_format.space_after = Pt(2)
            r = lp.add_run(f"({label})")
            r.font.name = "Malgun Gothic"
            r.font.size = Pt(10)
            r.font.bold = True
        ip = doc.add_paragraph()
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.space_after = Pt(4)
        ip.add_run().add_picture(str(fpath), width=Inches(width))

    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(8)
    cp.paragraph_format.space_after = Pt(2)
    r1 = cp.add_run(cap_title + "\n")
    r1.font.bold = True
    r1.font.name = "Malgun Gothic"
    r1.font.size = Pt(10)
    r2 = cp.add_run(cap_desc)
    r2.font.name = "Malgun Gothic"
    r2.font.size = Pt(9.5)
    if is_rev:
        r1.font.color.rgb = RED
        r2.font.color.rgb = RED

    src = doc.add_paragraph()
    src.paragraph_format.space_after = Pt(2)
    r = src.add_run("file: " + ", ".join("figures/" + f for _, f, _ in panels))
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(8)
    r.font.color.rgb = GRAY

out = os.path.join(OUT_DIR, "figures_v3.docx")
doc.save(out)
print(f"Saved {out}  figures={len(CAPTIONS)}  images={sum(len(v) for v in LAYOUT.values())}")
if missing:
    print("MISSING:", missing)
