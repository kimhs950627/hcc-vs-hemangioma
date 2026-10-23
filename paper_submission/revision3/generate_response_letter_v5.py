"""Build Response_to_Reviewers_v5.docx from the plain-text letter.

Same layout as generate_response_letter_v4.py (title 16pt, reviewer header 13pt blue,
question 9.5pt bold gray, answer 9.5pt). Edit the .txt, then rerun this script.
"""
import pathlib
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
SRC = HERE / "Response_to_Reviewers_v5.txt"
OUT = HERE / "Response_to_Reviewers_v5.docx"


def parse(lines):
    blocks = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].strip()
        if not line or set(line) <= {"-", "="}:
            i += 1
            continue
        if line.startswith("==="):
            blocks.append(("Title", line.strip("= ").strip()))
            i += 1
            continue
        if line.startswith("논문 제목") or line.startswith("저널"):
            meta = []
            while i < n and (lines[i].strip().startswith("논문 제목") or lines[i].strip().startswith("저널")):
                meta.append(lines[i].strip())
                i += 1
            blocks.append(("Meta", "\n".join(meta)))
            continue
        if line.startswith("●"):
            blocks.append(("Reviewer_Header", line))
            i += 1
            continue
        if re.match(r"^(주요 코멘트|마이너 코멘트|심사의견)", line):
            q = []
            while i < n and not lines[i].strip().startswith("[답변]"):
                if lines[i].strip() and not set(lines[i].strip()) <= {"-", "="}:
                    q.append(lines[i].strip())
                i += 1
            blocks.append(("Q", "\n".join(q)))
            continue
        if line.startswith("[답변]"):
            a = []
            while i < n:
                cur = lines[i].rstrip()
                if cur.strip().startswith("●") or re.match(r"^(주요 코멘트|마이너 코멘트|심사의견)", cur.strip()):
                    break
                if set(cur.strip()) <= {"-", "="} and cur.strip():
                    i += 1
                    break
                if cur.strip():
                    a.append(cur.strip())
                i += 1
            blocks.append(("A", "\n".join(a)))
            continue
        # intro / free paragraph or numbered list
        buf = []
        kind = "Intro" if not blocks or blocks[-1][0] != "Reviewer_Header" else "A"
        while i < n:
            cur = lines[i].rstrip()
            st = cur.strip()
            if not st or set(st) <= {"-", "="} or st.startswith("●") or re.match(r"^(주요 코멘트|마이너 코멘트|심사의견|\[답변\])", st):
                break
            buf.append(st)
            i += 1
        if buf:
            blocks.append((kind, "\n".join(buf)))
        else:
            i += 1
    return blocks


blocks = parse(SRC.read_text(encoding="utf-8").splitlines())

# merge consecutive Intro blocks into one paragraph
merged = []
for kind, text in blocks:
    if kind == "Intro" and merged and merged[-1][0] == "Intro":
        merged[-1] = ("Intro", merged[-1][1] + "\n" + text)
    elif kind == "Intro" and merged and merged[-1][0] == "A" and text.startswith("*"):
        merged[-1] = ("A", merged[-1][1] + "\n" + text)
    else:
        merged.append((kind, text))
blocks = merged

doc = docx.Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(1.0)
    s.left_margin = s.right_margin = Inches(1.0)

BLUE = RGBColor(0, 51, 153)
DARK_GRAY = RGBColor(80, 80, 80)

for kind, text in blocks:
    p = doc.add_paragraph()
    if kind == "Title":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.font.name = "Malgun Gothic"; r.font.size = Pt(16); r.font.bold = True
    elif kind == "Meta":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text); r.font.name = "Malgun Gothic"; r.font.size = Pt(10); r.font.italic = True
    elif kind == "Intro":
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(text); r.font.name = "Malgun Gothic"; r.font.size = Pt(10)
    elif kind == "Reviewer_Header":
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(text); r.font.name = "Malgun Gothic"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BLUE
    elif kind == "Q":
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(text); r.font.name = "Malgun Gothic"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = DARK_GRAY
    else:  # A
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(text); r.font.name = "Malgun Gothic"; r.font.size = Pt(9.5)

doc.save(str(OUT))
print("saved", OUT, "blocks", len(blocks))
for k, t in blocks[:6]:
    print(" ", k, t[:60].replace("\n", " / "))
