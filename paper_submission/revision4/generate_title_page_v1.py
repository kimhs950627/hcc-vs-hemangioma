# -*- coding: utf-8 -*-
"""Title page (표지) for KJFP submission — authors and affiliation supplied by
the corresponding author (2026-10-27).

Fields still needing author confirmation are marked 〔확인〕.
Run: python3 generate_title_page_v1.py
"""
import pathlib

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = pathlib.Path(__file__).resolve().parent

KO_TITLE = ("B-mode 복부 초음파에서 간세포암과 간혈관종 감별을 위한 딥러닝 모델의 "
            "Confidence Score와 임베딩 기반 유사도 점수: 단일기관 공개 데이터셋을 이용한 탐색적 연구")
EN_TITLE = ("Confidence Score and Embedding-Based Similarity Score of a Deep Learning Model for "
            "Differentiating Hepatocellular Carcinoma from Hemangioma on B-mode Abdominal Ultrasound: "
            "An Exploratory Study Using a Single-Center Public Dataset")
KO_RUNNING = "간암-혈관종 감별 이중 출력 지표"
EN_RUNNING = "Dual-output indices for HCC versus hemangioma"

AUTHORS_KO = "김현수¹, 이진희¹*"
AUTHORS_EN = "Hyunsu Kim¹, Jinhee Lee¹*"
AFFIL = ("¹ Presbyterian Medical Center (예수병원), Jeonju, Republic of Korea"
         " — Hyunsu Kim (ORCID 0009-0009-8896-5344), Jinhee Lee (ORCID 0000-0001-7762-4952)")

CORRESP = [
    ("교신저자", "Jinhee Lee (이진희), Presbyterian Medical Center (예수병원)"),
    ("주소", "전라북도 전주시 완산구 서원로 365 (중화산동1가), 54987, Republic of Korea 〔상세주소 확인〕"),
    ("E-mail", "jin56878@jesushospital.com"),
    ("전화 / 팩스", "〔전화번호〕 / 〔팩스번호〕"),
    ("ORCID", "0000-0001-7762-4952"),
]

EXTRA = [
    ("원고 관련 연락 (manuscript correspondence)", "Hyunsu Kim — kimhs950627@gmail.com"),
    ("연구비 (Funding)", "없음 (none)"),
    ("이해관계 (Conflicts of interest)", "없음 (none declared)"),
]

doc = docx.Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(3)

st = doc.styles['Normal']
st.font.name = 'Batang'
st.font.size = Pt(10)
st.paragraph_format.line_spacing = 2.0


def para(text, size=10, bold=False, align=None, space_before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    r = p.add_run(text)
    r.font.name = 'Batang'
    r.font.size = Pt(size)
    r.font.bold = bold
    return p


para("[표지 / Title page]", size=9, align=WD_ALIGN_PARAGRAPH.RIGHT)
para("국문 제목", bold=True, space_before=6)
para(KO_TITLE)
para("영문 제목 (English title)", bold=True, space_before=8)
para(EN_TITLE)

para("국문 소제목 (Running title, 30자 이내)", bold=True, space_before=8)
para(f"{KO_RUNNING}  [{len(KO_RUNNING)}자]")
para("영문 소제목 (Running title, 10 words 이내)", bold=True, space_before=8)
para(f"{EN_RUNNING}  [{len(EN_RUNNING.split())} words]")

para("저자 (Authors)", bold=True, space_before=8)
para(AUTHORS_KO)
para(AUTHORS_EN)
para("소속 및 ORCID", bold=True, space_before=8)
para(AFFIL)

para("교신저자 정보 (Corresponding author)", bold=True, space_before=8)
for label, value in CORRESP:
    para(f"- {label}: {value}")

para("기타", bold=True, space_before=8)
for label, value in EXTRA:
    para(f"- {label}: {value}")

para("", space_before=8)
para("* 교신저자 (corresponding author). 연구비 수혜 및 이해관계 관련 사항은 위와 같다.", size=9)

out = HERE / 'Title_Page_v1.docx'
doc.save(out)
print('wrote', out)
print('running title KO:', len(KO_RUNNING), 'chars /', KO_RUNNING)
print('running title EN:', len(EN_RUNNING.split()), 'words /', EN_RUNNING)
