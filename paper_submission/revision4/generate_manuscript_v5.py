# -*- coding: utf-8 -*-
"""manuscript_v5_source.md -> KJFP-formatted docx + plain text.

Formatting per KJFP rules: A4, 3 cm margins, double line spacing, Batang 10 pt,
page numbers at the bottom centre, in-text citations as superscripts with a
closing parenthesis (3+ consecutive numbers collapsed to a hyphen range).

Blue = sections changed in this round ([[B]] ... [[/B]] in the source).
Run: python3 generate_manuscript_v5.py
"""
import pathlib, re

import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
SRC = HERE / 'manuscript_v5_source.md'
BLUE = RGBColor(0, 0, 220)
BLACK = RGBColor(0, 0, 0)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin); run._r.append(instr); run._r.append(fld_end)
    run.font.name = 'Batang'
    run.font.size = Pt(10)


def compress(nums):
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append(f'{nums[i]}-{nums[j]}')
        else:
            out.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return ','.join(out)


CITE = re.compile(r'\[([0-9,\s]+)\]')


def add_rich_text(p, text, blue, size=Pt(10), bold=False, name='Batang'):
    pos = 0
    for m in CITE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            _style(r, blue, size, bold, name)
        nums = sorted({int(x) for x in re.findall(r'\d+', m.group(1))})
        r = p.add_run(compress(nums) + ')')
        _style(r, blue, size, bold, name)
        r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        _style(r, blue, size, bold, name)


def _style(run, blue, size, bold, name):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = BLUE if blue else BLACK
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), name)


doc = docx.Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(3)
add_page_number_footer(sec)

style = doc.styles['Normal']
style.font.name = 'Batang'
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Batang')
style.paragraph_format.line_spacing = 2.0
style.paragraph_format.space_after = Pt(0)

raw = SRC.read_text(encoding='utf-8')
blocks = [b for b in raw.split('\n') if b.strip()]

plain_lines = []
for block in blocks:
    line = block.strip()
    blue = False
    if line.startswith('[[B]]') and line.endswith('[[/B]]'):
        line = line[5:-6]
        blue = True
    plain_lines.append(CITE.sub(lambda m: compress(sorted({int(x) for x in re.findall(r'\d+', m.group(1))})) + ')', line))

    if line.startswith('### '):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 2.0
        add_rich_text(p, line[4:], blue, size=Pt(10), bold=True)
    elif line.startswith('## '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        add_rich_text(p, line[3:], blue, size=Pt(11), bold=True)
    elif line.startswith('# '):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(p, line[2:], blue, size=Pt(13), bold=True)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich_text(p, line, blue)

# ---- Figure legends (moved out of the tables file, KJFP order 8) ----
from captions_v5 import CAPTIONS as _FIGCAPS

_p = doc.add_paragraph()
_p.paragraph_format.space_before = Pt(18)
_r = _p.add_run('그림 설명문 (Figure Legends)')
_r.font.name = 'Batang'
_r.font.size = Pt(11)
_r.font.bold = True
plain_lines.append('## 그림 설명문 (Figure Legends)')

for _title, _desc, _ in _FIGCAPS:
    _t = doc.add_paragraph()
    _t.paragraph_format.space_before = Pt(8)
    _t.paragraph_format.line_spacing = 1.5
    _rt = _t.add_run(_title)
    _rt.font.name = 'Batang'
    _rt.font.size = Pt(10)
    _rt.font.bold = True
    _d = doc.add_paragraph()
    _d.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _rd = _d.add_run(_desc)
    _rd.font.name = 'Batang'
    _rd.font.size = Pt(10)
    plain_lines.append(_title)
    plain_lines.append(_desc)

out_docx = HERE / 'manuscript_v5_kjfp.docx'
doc.save(out_docx)
(HERE / 'manuscript_v5_plain.txt').write_text('\n\n'.join(plain_lines), encoding='utf-8')
print('wrote', out_docx)
print('wrote', HERE / 'manuscript_v5_plain.txt')
